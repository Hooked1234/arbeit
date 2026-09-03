# -*- coding: utf-8 -*-
"""Baut die Lehrendenfassung (DOCX) aus content.py."""
import os, sys
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

import content as K

NAVY = RGBColor(0x17, 0x3B, 0x57)
GREY = RGBColor(0x55, 0x5F, 0x66)
LIGHT = "EEF2F5"
BAND = "173B57"
FONT = "Calibri"

OUT = sys.argv[1]
DIA = sys.argv[2]
A4_W, A4_H = Cm(21.0), Cm(29.7)
A3_W, A3_H = Cm(29.7), Cm(42.0)
MARGIN = Cm(1.8)


# ------------------------------------------------------------------ Helfer
def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, bold=False, size=9, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color or NAVY


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def drop_trailing_empty(doc):
    """Leere Absaetze am Ende entfernen, damit Abschnittswechsel keine
    Leerseite erzeugen."""
    body = doc.element.body
    while True:
        kids = [c for c in body.iterchildren()
                if c.tag != qn("w:sectPr")]
        if not kids:
            return
        last = kids[-1]
        if last.tag == qn("w:p") and not "".join(last.itertext()).strip():
            body.remove(last)
        else:
            return


def table(doc, rows, widths, header=True, size=9):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, r in enumerate(rows):
        row = t.add_row()
        for j, val in enumerate(r):
            c = row.cells[j]
            c.width = widths[j]
            cell_text(c, val, bold=(header and i == 0), size=size,
                      color=NAVY if not (header and i == 0) else RGBColor(
                          0xFF, 0xFF, 0xFF))
            if header and i == 0:
                shade(c, BAND)
            elif i % 2 == 0:
                shade(c, LIGHT)
    for i, row in enumerate(t.rows):
        no_split(row)
        for j, c in enumerate(row.cells):
            c.width = widths[j]
    if header and t.rows:
        repeat_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def para(doc, text, size=10, bold=False, italic=False, color=None,
         space_after=6, space_before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color or NAVY
    return p


def heading(doc, text, level=1):
    sizes = {0: 20, 1: 14, 2: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.style = doc.styles["Heading %d" % max(1, level)]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(sizes[level])
    r.bold = True
    r.font.color.rgb = NAVY
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "8")
        bot.set(qn("w:space"), "3")
        bot.set(qn("w:color"), BAND)
        b.append(bot)
        pPr.append(b)
    return p


def bullets(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY


def numbers(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY


def callout(doc, label, text, size=10):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = Cm(17.4)
    shade(c, LIGHT)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + "  ")
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = True
    r.font.color.rgb = NAVY
    r2 = p.add_run(text)
    r2.font.name = FONT
    r2.font.size = Pt(size)
    r2.font.color.rgb = NAVY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def landscape_a3(doc):
    drop_trailing_empty(doc)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = A3_H, A3_W
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, a, Cm(1.2))
    return s


def portrait_a4(doc):
    drop_trailing_empty(doc)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = A4_W, A4_H
    s.left_margin = s.right_margin = MARGIN
    s.top_margin = s.bottom_margin = Cm(1.6)
    return s


# -------------------------------------------------------------------- Bau
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.font.color.rgb = NAVY

s = doc.sections[0]
s.page_width, s.page_height = A4_W, A4_H
s.left_margin = s.right_margin = MARGIN
s.top_margin = s.bottom_margin = Cm(1.6)

para(doc, "LEHRENDENFASSUNG", size=9, bold=True, color=GREY, space_after=2)
para(doc, "UML-Klassendiagramme modellieren", size=22, bold=True,
     space_after=2)
para(doc, "Musterlösungen, Moderation und Bewertungsraster für drei "
          "Kurzübungen à 20 Minuten", size=11, color=GREY, space_after=10)
table(doc, [["3 Aufgaben", "60 Punkte", "3 × 20 Minuten",
             "Stand " + K.STAND]],
      [Cm(4.35)] * 4, header=True, size=10)

heading(doc, "1  Durchführung")
para(doc, "Die Aufgabenfolien enthalten ausschließlich Titel und Szenariotext. "
          "Alles, was modelliert werden soll, steht im Text; die Notation "
          "kommt aus der Theoriepräsentation des Dozenten.")
para(doc, "Ablauf je Aufgabe", bold=True, size=10, space_before=6,
     space_after=3)
table(doc, [["Zeit", "Was geschieht", "Wer"]] + [list(r) for r in K.ABLAUF],
      [Cm(2.6), Cm(12.4), Cm(2.4)])
para(doc, "Priorisierung, falls die Zeit knapp wird", bold=True, size=10,
     space_before=4, space_after=3)
para(doc, "Ab Minute 10 ansagen, in dieser Reihenfolge zu vervollständigen:",
     size=10, space_after=3)
numbers(doc, K.PRIORITAET)

heading(doc, "2  Didaktische Leitplanken")
bullets(doc, K.LEITPLANKEN)

heading(doc, "3  Gesamtraster")
para(doc, "Je Aufgabe 20 Punkte, insgesamt 60. Die Gewichtung ist in allen "
          "drei Aufgaben identisch.", space_after=4)
table(doc, [["Dimension", "Gewichtung", "Leitfrage"]] +
      [list(r) for r in K.GESAMTRASTER],
      [Cm(5.0), Cm(2.4), Cm(10.0)])

W1, W2 = Cm(8.2), Cm(9.2)
for idx, A in enumerate(K.AUFGABEN):
    portrait_a4(doc)
    para(doc, "MUSTERLÖSUNG %d" % A["nr"], size=9, bold=True, color=GREY,
         space_after=2)
    heading(doc, "%d  %s" % (3 + A["nr"], A["titel"]))
    table(doc, [["Umfang", "Schwerpunkt"], [A["umfang"], A["schwerpunkt"]]],
          [Cm(6.0), Cm(11.4)])

    para(doc, "Szenariotext (identisch zur Aufgabenfolie)", bold=True,
         size=10, space_before=4, space_after=3)
    callout(doc, "", A["szenario"], size=9)

    para(doc, "Moderation", bold=True, size=11, space_before=6, space_after=3)
    callout(doc, "Minute 5 – offen fragen:", A["hinweis5"])
    callout(doc, "Minute 10 – eingrenzen:", A["hinweis10"])
    callout(doc, "Checkpoint-Satz zur Auflösung:", A["checkpoint"])

    para(doc, "Beziehungsübersicht", bold=True, size=11, space_before=6,
         space_after=3)
    table(doc, [["Beziehung", "Typ", "Multiplizität", "Bedeutung"]] +
          [list(r) for r in A["beziehungen"]],
          [Cm(5.4), Cm(3.6), Cm(2.6), Cm(5.8)], size=8.5)

    para(doc, "Prüfinstanzen", bold=True, size=11, space_before=6,
         space_after=3)
    table(doc, [["Fall", "Konstellation", "Erwartung"]] +
          [list(r) for r in A["pruef"]],
          [Cm(3.4), Cm(7.0), Cm(7.0)], size=8.5)

    para(doc, "Typische Fehler", bold=True, size=11, space_before=6,
         space_after=3)
    bullets(doc, A["fehler"], size=9.5)

    para(doc, "Zulässige Alternativen", bold=True, size=11, space_before=6,
         space_after=3)
    bullets(doc, A["alternativen"], size=9.5)

    para(doc, "Bewertungsraster (20 Punkte)", bold=True, size=11,
         space_before=6, space_after=3)
    table(doc, [["Kriterium", "Punkte", "Voll erfüllt, wenn …"]] +
          [list(r) for r in A["raster"]],
          [Cm(5.0), Cm(1.8), Cm(10.6)], size=8.5)

    para(doc, "Rückführbarkeit: jeder Satz, jedes Modellelement", bold=True,
         size=11, space_before=6, space_after=3)
    para(doc, "Ein Modellelement ohne Satz ist überflüssig, ein Satz ohne "
              "Modellelement ist ein Fehler in der Aufgabe.", size=9,
         italic=True, color=GREY, space_after=4)
    table(doc, [["Satz im Szenario", "Modellelement"]] +
          [list(r) for r in A["rueck"]], [W1, W2], size=8.5)

    # Diagramm auf eigener A3-Querseite
    landscape_a3(doc)
    para(doc, "Musterlösung Aufgabe %d – %s" % (A["nr"], A["titel"]),
         size=13, bold=True, space_after=4)
    png = os.path.join(DIA, {
        1: "Aufgabe_1_Bestellsystem_eines_Onlineshops.png",
        2: "Aufgabe_2_Hotelverwaltung.png",
        3: "Aufgabe_3_Veranstaltungsmanagement.png"}[A["nr"]])
    doc.add_picture(png, width=Cm(39.0))
    para(doc, "Die Vektorfassung liegt separat als SVG und als "
              "diagrams.net-Quelle vor.", size=9, italic=True, color=GREY,
         space_before=4)

portrait_a4(doc)
heading(doc, "7  Abnahme vor der Veranstaltung")
bullets(doc, K.ABNAHME)
para(doc, "Solange der Zeittest aussteht, ist die Freigabe offen. Die "
          "Zeitangaben beruhen auf Schätzung, nicht auf Messung.",
     size=10, italic=True, color=GREY, space_before=6)

doc.save(OUT)
print("geschrieben:", OUT)
