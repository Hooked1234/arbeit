# -*- coding: utf-8 -*-
"""Druckfassung der Aufgaben: identischer Text wie die Folien, nichts sonst."""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as K

NAVY = RGBColor(0x17, 0x3B, 0x57)
GREY = RGBColor(0x55, 0x5F, 0x66)
BAND = "173B57"
FONT = "Calibri"
OUT = sys.argv[1]

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.font.color.rgb = NAVY

s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.2)
s.top_margin = s.bottom_margin = Cm(1.8)


def line(text, size=11, bold=False, italic=False, color=None, after=6,
         before=0, rule=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color or NAVY
    if rule:
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "8")
        bot.set(qn("w:space"), "4")
        bot.set(qn("w:color"), BAND)
        b.append(bot)
        pPr.append(b)
    return p


def keep_with_next(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pPr.append(el)


line("UML-Klassendiagramme", size=9, bold=True, color=GREY, after=2)
line("Übungsaufgaben", size=20, bold=True, after=2)
line("Modellierung betrieblicher Informationssysteme · 12.08.2026 · "
     "drei Aufgaben à 20 Minuten", size=10, color=GREY, after=16, rule=True)

for i, A in enumerate(K.AUFGABEN):
    h = line("Aufgabe %d – %s" % (A["nr"], A["titel"]), size=14, bold=True,
             before=(0 if i == 0 else 20), after=6)
    keep_with_next(h)
    p = line(A["szenario"], size=11, after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.save(OUT)
print("geschrieben:", OUT)
