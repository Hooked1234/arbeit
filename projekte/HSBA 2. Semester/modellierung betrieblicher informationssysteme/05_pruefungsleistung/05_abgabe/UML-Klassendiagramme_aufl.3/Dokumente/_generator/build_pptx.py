# -*- coding: utf-8 -*-
"""Aufgabenfolien im Format des Dozenten.

Basis ist seine Vorlage `Aufgaben_EPK_Keine_Loesung.pptx`: Master, Theme und
Layout bleiben unveraendert, Folien 4 und 5 entfallen, die verbleibenden drei
werden mit unseren Szenarien belegt. Dokumenteigenschaften werden neu gesetzt,
damit die Datei nicht als seine ausgewiesen wird.
"""
import copy, shutil, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt


def kein_aufzaehlungszeichen(para):
    """Aufzaehlungszeichen aus dem Layout unterdruecken und Einzug nullen."""
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))

import content as K

TEMPLATE = sys.argv[1]
OUT = sys.argv[2]
FS = {1: Pt(18), 2: Pt(18), 3: Pt(18)}   # wie die Reklamationsfolie

shutil.copyfile(TEMPLATE, OUT)
p = Presentation(OUT)

# Folien auf drei kuerzen
ids = p.slides._sldIdLst
for sid in list(ids)[3:]:
    rId = sid.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/"
        "relationships}id")
    p.part.drop_rel(rId)
    ids.remove(sid)

# Volle Breite wie auf der Reklamationsfolie des Dozenten
FULL_L, FULL_W = Cm(0.7), Cm(32.6)

for i, (slide, A) in enumerate(zip(p.slides, K.AUFGABEN)):
    title = None
    body = None
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.placeholder_format is not None and sh.placeholder_format.idx == 0:
            title = sh
        else:
            body = sh
    title.text_frame.text = "Aufgabe %d - %s" % (A["nr"], A["titel"])

    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    body.left, body.width = FULL_L, FULL_W
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = A["szenario"]
    run.font.size = FS[A["nr"]]
    para.line_spacing = 1.0
    para.space_after = Pt(0)
    kein_aufzaehlungszeichen(para)

cp = p.core_properties
cp.title = "UML-Klassendiagramme - Uebungsaufgaben"
cp.subject = "MOBIS, Pruefungsleistung 12.08.2026"
cp.author = ""
cp.last_modified_by = ""
cp.comments = ("Aufgabenfolien nach dem Format von Aufgabe 5 - Reklamation. "
               "Vorlage, Master und Theme stammen aus dem Kursmaterial.")
p.save(OUT)
print("geschrieben:", OUT)
for i, s in enumerate(p.slides):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text:
            print("  F%d %-8s %.1f x %.1f cm | %d Woerter"
                  % (i + 1, "Titel" if sh.width > Cm(30) and
                     sh.height < Cm(4) else "Text",
                     sh.width / 360000, sh.height / 360000,
                     len(sh.text_frame.text.split())))
