# -*- coding: utf-8 -*-
"""Aufgabenblaetter zum Austeilen.

Blatt 1: Szenariotext und Arbeitsauftrag zu Aufgabe 1.
Blatt 2: Auftrag, Protokollbogen und Schlussfrage zu Aufgabe 2.

Der Text ist identisch zu den Aufgabenfolien. Bewusst nicht enthalten:
Pflichtklassenlisten, Beziehungstypen, Notationsregeln - sie naehmen die
Loesung vorweg. Die eigene Runde vom 12.08. hatte eine solche Checkliste nur
bei der leichtesten Aufgabe; fuer eine Vertiefung ist sie zu viel.

Aufruf:  python _generator/build_docx_aufgaben.py <ziel.docx>
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as K

NAVY = RGBColor(0x17, 0x3B, 0x57)
GREY = RGBColor(0x55, 0x5F, 0x66)
BAND = "173B57"
LIGHT = "EEF2F5"
FONT = "Calibri"
OUT = sys.argv[1]

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.font.color.rgb = NAVY

s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
s.top_margin = s.bottom_margin = Cm(1.6)


def line(text, size=11, bold=False, italic=False, color=None, after=6,
         before=0, rule=False, justify=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color or NAVY
    if rule:
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "8")
        bot.set(qn("w:space"), "4")
        bot.set(qn("w:color"), BAND)
        b.append(bot)
        pPr.append(b)
    return p


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color or NAVY


def kopfzeile():
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for c, w, txt in ((t.rows[0].cells[0], Cm(8.5), "Gruppe"),
                      (t.rows[0].cells[1], Cm(8.5), "Namen")):
        c.width = w
        cell_text(c, txt + ": ", bold=True, size=9, color=GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def protokollbogen(zeilen=10):
    kopf = ["Nr.", "Stelle im Diagramm", "Betroffene Aussage im Text",
            "Fehler oder zulaessig", "Korrektur"]
    breiten = [Cm(0.9), Cm(3.9), Cm(4.6), Cm(2.6), Cm(4.9)]
    t = doc.add_table(rows=zeilen + 1, cols=5)
    t.style = "Table Grid"
    t.autofit = False
    for j, (txt, w) in enumerate(zip(kopf, breiten)):
        c = t.rows[0].cells[j]
        c.width = w
        shade(c, LIGHT)
        cell_text(c, txt, bold=True, size=8.5, color=GREY)
    for i in range(1, zeilen + 1):
        for j, w in enumerate(breiten):
            c = t.rows[i].cells[j]
            c.width = w
            cell_text(c, str(i) if j == 0 else "", size=9)
        t.rows[i].height = Cm(0.95)
    return t


def linien(anzahl=4):
    for _ in range(anzahl):
        line("", after=2, rule=True)


# ------------------------------------------------------------------ Blatt 1
A = K.A1
line("UML-Klassendiagramm - Vertiefung", size=9, bold=True, color=GREY,
     after=2)
line("Aufgabe %d: %s" % (A["nr"], A["titel"]), size=19, bold=True, after=2)
line("Modellierung betrieblicher Informationssysteme  |  %s  |  45 Minuten"
     % K.TERMIN, size=10, color=GREY, after=12, rule=True)
kopfzeile()

for abs_ in A["szenario"].split("\n\n"):
    line(abs_, size=11, after=8, justify=True)

line("Arbeitsauftrag", size=12, bold=True, before=8, after=4)
for punkt in A["auftrag"]:
    p = line("-  " + punkt, size=10.5, after=3)
    p.paragraph_format.left_indent = Cm(0.4)

# ------------------------------------------------------------------ Blatt 2
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

B = K.A2
line("UML-Klassendiagramm - Vertiefung", size=9, bold=True, color=GREY,
     after=2)
line("Aufgabe %d: %s" % (B["nr"], B["titel"]), size=19, bold=True, after=2)
line("Der Entwurf liegt als eigenes Blatt bei  |  45 Minuten", size=10,
     color=GREY, after=12, rule=True)
kopfzeile()

for abs_ in B["szenario"].split("\n\n"):
    line(abs_, size=11, after=8, justify=True)

line("Zu beachten", size=12, bold=True, before=6, after=4)
for punkt in B["auftrag"]:
    p = line("-  " + punkt, size=10.5, after=3)
    p.paragraph_format.left_indent = Cm(0.4)

line("Protokoll", size=12, bold=True, before=10, after=5)
protokollbogen(10)

doc.add_paragraph().paragraph_format.space_after = Pt(10)
line("Schlussfrage", size=12, bold=True, before=6, after=4)
line("Welchen der gefundenen Fehler haette das Werkzeug nicht machen koennen, "
     "wenn der Szenariotext an einer Stelle genauer gewesen waere?", size=10.5,
     after=8, justify=True)
linien(4)

doc.save(OUT)
print("geschrieben:", OUT)
