# -*- coding: utf-8 -*-
"""Baut die Lehrendenfassung der Vertiefung (DOCX) aus content.py.

Aufruf:  python _generator/build_docx.py <ziel.docx> <Diagrammordner>
"""
import os, sys
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

import content as K

NAVY = RGBColor(0x17, 0x3B, 0x57)
GREY = RGBColor(0x55, 0x5F, 0x66)
LIGHT = "EEF2F5"
BAND = "173B57"
FONT = "Calibri"

OUT = sys.argv[1]
DIA = sys.argv[2]
A4_W, A4_H = Cm(21.0), Cm(29.7)
A3_W, A3_H = Cm(29.7), Cm(42.0)
MARGIN = Cm(1.8)


# ------------------------------------------------------------------ Helfer
def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, bold=False, size=9, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color or NAVY


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def drop_trailing_empty(doc):
    """Leere Absaetze am Ende entfernen, damit Abschnittswechsel keine
    Leerseite erzeugen."""
    body = doc.element.body
    while True:
        kids = [c for c in body.iterchildren()
                if c.tag != qn("w:sectPr")]
        if not kids:
            return
        last = kids[-1]
        if last.tag == qn("w:p") and not "".join(last.itertext()).strip():
            body.remove(last)
        else:
            return


def table(doc, rows, widths, header=True, size=9):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, r in enumerate(rows):
        row = t.add_row()
        for j, val in enumerate(r):
            c = row.cells[j]
            c.width = widths[j]
            cell_text(c, val, bold=(header and i == 0), size=size,
                      color=NAVY if not (header and i == 0) else RGBColor(
                          0xFF, 0xFF, 0xFF))
            if header and i == 0:
                shade(c, BAND)
            elif i % 2 == 0:
                shade(c, LIGHT)
    for i, row in enumerate(t.rows):
        no_split(row)
        for j, c in enumerate(row.cells):
            c.width = widths[j]
    if header and t.rows:
        repeat_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def para(doc, text, size=10, bold=False, italic=False, color=None,
         space_after=6, space_before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color or NAVY
    return p


def heading(doc, text, level=1):
    sizes = {0: 20, 1: 14, 2: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.style = doc.styles["Heading %d" % max(1, level)]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(sizes[level])
    r.bold = True
    r.font.color.rgb = NAVY
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "8")
        bot.set(qn("w:space"), "3")
        bot.set(qn("w:color"), BAND)
        b.append(bot)
        pPr.append(b)
    return p


def bullets(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY


def numbers(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY


def callout(doc, label, text, size=10):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = Cm(17.4)
    shade(c, LIGHT)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + "  ")
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = True
    r.font.color.rgb = NAVY
    r2 = p.add_run(text)
    r2.font.name = FONT
    r2.font.size = Pt(size)
    r2.font.color.rgb = NAVY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def landscape_a3(doc):
    drop_trailing_empty(doc)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = A3_H, A3_W
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, a, Cm(1.2))
    return s


def portrait_a4(doc):
    drop_trailing_empty(doc)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = A4_W, A4_H
    s.left_margin = s.right_margin = MARGIN
    s.top_margin = s.bottom_margin = Cm(1.6)
    return s


# ------------------------------------------------------------- Zusatzhelfer
def textbox(doc, text, size=9):
    """Mehrabsaetziger Kasten fuer den ausgeteilten Wortlaut."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = Cm(17.4)
    shade(c, LIGHT)
    c.text = ""
    first = True
    for abs_ in text.split("\n\n"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(abs_)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# -------------------------------------------------------------------- Bau
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.font.color.rgb = NAVY

s = doc.sections[0]
s.page_width, s.page_height = A4_W, A4_H
s.left_margin = s.right_margin = MARGIN
s.top_margin = s.bottom_margin = Cm(1.6)

para(doc, "LEHRENDENFASSUNG", size=9, bold=True, color=GREY, space_after=2)
para(doc, "UML-Klassendiagramm - Vertiefung", size=22, bold=True,
     space_after=2)
para(doc, "Musterloesungen, Moderation und Bewertungsraster fuer zwei "
          "Uebungen a 45 Minuten", size=11, color=GREY, space_after=10)
table(doc, [["2 Aufgaben", "40 Punkte", "2 x 45 Minuten",
             "Termin " + K.TERMIN]],
      [Cm(4.35)] * 4, header=True, size=10)

heading(doc, "1  Durchfuehrung")
para(doc, "Aufgabe 1 wird modelliert, Aufgabe 2 prueft einen fremden Entwurf "
          "zum selben Text. Die Musterloesung zu Aufgabe 1 wird am Ende von "
          "Block 1 vorgestellt und danach wieder eingesammelt: In Block 2 ist "
          "der Szenariotext der Massstab, nicht das eigene Diagramm.")

para(doc, "Ablauf Aufgabe 1", bold=True, size=10, space_before=6,
     space_after=3)
table(doc, [["Zeit", "Was geschieht", "Wer"]] + [list(r) for r in K.ABLAUF_1],
      [Cm(2.6), Cm(12.4), Cm(2.4)])

para(doc, "Ablauf Aufgabe 2", bold=True, size=10, space_before=6,
     space_after=3)
table(doc, [["Zeit", "Was geschieht", "Wer"]] + [list(r) for r in K.ABLAUF_2],
      [Cm(2.6), Cm(12.4), Cm(2.4)])

para(doc, "Priorisierung, falls die Zeit in Aufgabe 1 knapp wird", bold=True,
     size=10, space_before=4, space_after=3)
para(doc, "Ab Minute 20 ansagen, in dieser Reihenfolge zu vervollstaendigen:",
     size=10, space_after=3)
numbers(doc, K.PRIORITAET)

heading(doc, "2  Didaktische Leitplanken")
bullets(doc, K.LEITPLANKEN)

# ----------------------------------------------------------------- Aufgabe 1
portrait_a4(doc)
A = K.A1
para(doc, "AUFGABE 1 - MUSTERLOESUNG", size=9, bold=True, color=GREY,
     space_after=2)
heading(doc, "3  %s" % A["titel"])
table(doc, [["Art", "Umfang", "Schwerpunkt"],
            [A["art"], A["umfang"], A["schwerpunkt"]]],
      [Cm(2.4), Cm(5.0), Cm(10.0)])

para(doc, "Szenariotext, identisch zur Aufgabenfolie", bold=True, size=10,
     space_before=6, space_after=3)
textbox(doc, A["szenario"])

para(doc, "Arbeitsauftrag auf dem Aufgabenblatt", bold=True, size=10,
     space_before=4, space_after=3)
bullets(doc, A["auftrag"], size=9.5)

para(doc, "Moderation", bold=True, size=11, space_before=6, space_after=3)
callout(doc, "Minute 10, offen gestellt:", A["hinweis1"])
callout(doc, "Minute 20, eingrenzend:", A["hinweis2"])
callout(doc, "Checkpoint-Satz zur Aufloesung:", A["checkpoint"])

para(doc, "Beziehungsuebersicht", bold=True, size=11, space_before=6,
     space_after=3)
table(doc, [["Beziehung", "Typ", "Kardinalitaet", "Bedeutung"]] +
      [list(r) for r in A["beziehungen"]],
      [Cm(5.4), Cm(3.6), Cm(2.6), Cm(5.8)], size=8.5)

para(doc, "Pruefinstanzen", bold=True, size=11, space_before=6, space_after=3)
table(doc, [["Fall", "Konstellation", "Erwartung"]] +
      [list(r) for r in A["pruef"]],
      [Cm(3.4), Cm(7.0), Cm(7.0)], size=8.5)

para(doc, "Typische Fehler", bold=True, size=11, space_before=6, space_after=3)
bullets(doc, A["fehler"], size=9.5)

para(doc, "Zulaessige Alternativen", bold=True, size=11, space_before=6,
     space_after=3)
bullets(doc, A["alternativen"], size=9.5)

para(doc, "Bewertungsraster, 20 Punkte", bold=True, size=11, space_before=6,
     space_after=3)
table(doc, [["Kriterium", "Punkte", "Voll erfuellt, wenn ..."]] +
      [list(r) for r in A["raster"]],
      [Cm(5.0), Cm(1.8), Cm(10.6)], size=8.5)

para(doc, "Rueckfuehrbarkeit: jeder Satz, jedes Modellelement", bold=True,
     size=11, space_before=6, space_after=3)
para(doc, "Ein Modellelement ohne Satz ist ueberfluessig, ein Satz ohne "
          "Modellelement ist ein Fehler in der Aufgabe.", size=9, italic=True,
     color=GREY, space_after=4)
table(doc, [["Satz im Szenario", "Modellelement"]] +
      [list(r) for r in A["rueck"]], [Cm(8.2), Cm(9.2)], size=8.5)

landscape_a3(doc)
para(doc, "Musterloesung Aufgabe 1 - %s" % A["titel"], size=13, bold=True,
     space_after=4)
doc.add_picture(os.path.join(DIA, "Stufe_1_Ist-Modell.png"), width=Cm(38.5))
para(doc, "Wird am Ende von Block 1 gezeigt und danach wieder eingesammelt. "
          "Vektorfassung als SVG und als diagrams.net-Quelle daneben.",
     size=9, italic=True, color=GREY, space_before=4)

# ----------------------------------------------------------------- Aufgabe 2
portrait_a4(doc)
B = K.A2
para(doc, "AUFGABE 2 - MUSTERLOESUNG", size=9, bold=True, color=GREY,
     space_after=2)
heading(doc, "4  %s" % B["titel"])
table(doc, [["Art", "Umfang", "Schwerpunkt"],
            [B["art"], B["umfang"], B["schwerpunkt"]]],
      [Cm(2.4), Cm(5.0), Cm(10.0)])

para(doc, "Auftragstext, identisch zur Aufgabenfolie", bold=True, size=10,
     space_before=6, space_after=3)
textbox(doc, B["szenario"])

para(doc, "Hinweise auf dem Aufgabenblatt", bold=True, size=10, space_before=4,
     space_after=3)
bullets(doc, B["auftrag"], size=9.5)

para(doc, "Herkunft des Entwurfs", bold=True, size=11, space_before=6,
     space_after=3)
para(doc, "Der Entwurf stammt aus einem KI-Werkzeug, ist aber fuer den "
          "Unterricht kuratiert: Zahl und Art der Fehler sind bewusst so "
          "gewaehlt, dass jeder Befund einen eigenen Lernpunkt trifft. Auf dem "
          "Aufgabenblatt steht deshalb, dass ein KI-Werkzeug den Entwurf "
          "erzeugt hat - nicht, dass es ein unveraenderter Output eines "
          "bestimmten Produkts sei. Wer danach fragt, bekommt diese Antwort.",
     size=9.5)

para(doc, "Moderation", bold=True, size=11, space_before=6, space_after=3)
callout(doc, "Minute 12, offen gestellt:", B["hinweis1"])
callout(doc, "Minute 20, eingrenzend:", B["hinweis2"])
callout(doc, "Checkpoint-Satz zur Aufloesung:", B["checkpoint"])

para(doc, "Fundliste - das ist die Musterloesung", bold=True, size=11,
     space_before=6, space_after=3)
para(doc, "Acht Fehler, zwei zulaessige Abweichungen. Volle Punktzahl bei "
          "sechs gefundenen Fehlern und mindestens einer richtig als zulaessig "
          "eingestuften Abweichung.", size=9, italic=True, color=GREY,
     space_after=4)
table(doc, [["Nr.", "Stelle im Entwurf", "Aussage im Text", "Befund",
             "Korrektur"]] + [list(r) for r in B["fundliste"]],
      [Cm(1.0), Cm(4.4), Cm(5.0), Cm(2.2), Cm(4.8)], size=8.5)

para(doc, "Antwort auf die Schlussfrage", bold=True, size=11, space_before=6,
     space_after=3)
callout(doc, "", B["schlussantwort"], size=9.5)

para(doc, "Typische Fehler", bold=True, size=11, space_before=6, space_after=3)
bullets(doc, B["fehler"], size=9.5)

para(doc, "Zulaessige Alternativen", bold=True, size=11, space_before=6,
     space_after=3)
bullets(doc, B["alternativen"], size=9.5)

para(doc, "Bewertungsraster, 20 Punkte", bold=True, size=11, space_before=6,
     space_after=3)
table(doc, [["Kriterium", "Punkte", "Voll erfuellt, wenn ..."]] +
      [list(r) for r in B["raster"]],
      [Cm(5.0), Cm(1.8), Cm(10.6)], size=8.5)

landscape_a3(doc)
para(doc, "Aufgabe 2 - der auszuteilende KI-Entwurf mit den zehn Befunden",
     size=13, bold=True, space_after=4)
doc.add_picture(os.path.join(DIA, "Aufgabe_2_KI-Entwurf.png"), width=Cm(38.5))
para(doc, "Zum Vergleich daneben die Musterloesung aus Aufgabe 1 legen. Die "
          "Nummern der Fundliste beziehen sich auf dieses Blatt.",
     size=9, italic=True, color=GREY, space_before=4)

# ---------------------------------------------------------------- Betreuung
portrait_a4(doc)
heading(doc, "5  Betreuung zu zweit")
para(doc, "Die Aufteilung dokumentiert zugleich den Individualanteil "
          "innerhalb der Gruppenleistung.", space_after=4)
table(doc, [["Rolle", "Person A", "Person B"]] +
      [list(r) for r in K.ROLLEN], [Cm(5.8), Cm(5.8), Cm(5.8)], size=9)

para(doc, "Hinweisstaffel - woertlich geben", bold=True, size=11,
     space_before=8, space_after=3)
para(doc, "Beide Betreuenden sagen denselben Satz. Ausserhalb der Staffel "
          "werden keine inhaltlichen Hinweise gegeben, damit keine Gruppe "
          "einen Vorteil hat.", size=9, italic=True, color=GREY, space_after=4)
table(doc, [["Zeitpunkt", "Wortlaut"],
            ["Aufgabe 1, Minute 10", K.A1["hinweis1"]],
            ["Aufgabe 1, Minute 20", K.A1["hinweis2"]],
            ["Aufgabe 2, Minute 12", K.A2["hinweis1"]],
            ["Aufgabe 2, Minute 20", K.A2["hinweis2"]]],
      [Cm(4.4), Cm(13.0)], size=9)

para(doc, "Antwortkatalog fuer Rueckfragen", bold=True, size=11,
     space_before=8, space_after=3)
para(doc, "Abgestimmte Antworten. Was hier nicht steht, wird mit einer "
          "Gegenfrage an den Text zurueckgegeben.", size=9, italic=True,
     color=GREY, space_after=4)
table(doc, [["Frage", "Antwort"]] + [list(r) for r in K.FAQ],
      [Cm(7.0), Cm(10.4)], size=8.5)

heading(doc, "6  Abnahme vor der Veranstaltung")
bullets(doc, K.ABNAHME)
para(doc, "Solange der Zeittest aussteht, ist die Freigabe offen. Die "
          "Zeitangaben beruhen auf Schaetzung, nicht auf Messung - und "
          "Aufgabe 2 ist eine Aufgabenart ohne eigenen Erfahrungswert.",
     size=10, italic=True, color=GREY, space_before=6)

doc.save(OUT)
print("geschrieben:", OUT)
