# -*- coding: utf-8 -*-
"""Aufgabenfolien im Format des Dozenten, zwei Folien.

Basis ist seine Vorlage `Aufgaben_EPK_Keine_Loesung.pptx`: Master, Theme und
Layout bleiben unveraendert, die Folien 3 bis 5 entfallen, die verbleibenden
zwei werden mit unseren Texten belegt. Dokumenteigenschaften werden neu
gesetzt, damit die Datei nicht als seine ausgewiesen wird.

Abweichung von der Vorlage, bewusst: Die Reklamationsfolie des Dozenten traegt
rund 175 Woerter bei 18 pt. Der Text zu Aufgabe 1 hat rund 360 Woerter, weil
die Vertiefung ein groesseres Modell traegt. Die Schriftgroesse geht deshalb
auf 14 pt herunter. Massgeblich zum Mitlesen ist ohnehin das Aufgabenblatt,
das gleichzeitig ausgeteilt wird; die Folie dient dem Vorlesen und der
Projektion.

Aufruf:  python _generator/build_pptx.py <Dozentenvorlage.pptx> <ziel.pptx>
"""
import shutil
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

import content as K

TEMPLATE = sys.argv[1]
OUT = sys.argv[2]

FS = {1: Pt(14), 2: Pt(16)}
FULL_L, FULL_W = Cm(0.7), Cm(32.6)


def kein_aufzaehlungszeichen(para):
    """Aufzaehlungszeichen aus dem Layout unterdruecken und Einzug nullen."""
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


shutil.copyfile(TEMPLATE, OUT)
p = Presentation(OUT)

# Auf zwei Folien kuerzen
ids = p.slides._sldIdLst
for sid in list(ids)[2:]:
    rId = sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/"
                  "relationships}id")
    p.part.drop_rel(rId)
    ids.remove(sid)

for slide, A in zip(p.slides, K.AUFGABEN):
    title = body = None
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.placeholder_format is not None and sh.placeholder_format.idx == 0:
            title = sh
        else:
            body = sh
    title.text_frame.text = "Aufgabe %d - %s" % (A["nr"], A["titel"])

    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    body.left, body.width = FULL_L, FULL_W

    absaetze = A["szenario"].split("\n\n")
    for i, txt in enumerate(absaetze):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = txt
        run.font.size = FS[A["nr"]]
        para.line_spacing = 1.0
        para.space_after = Pt(6)
        kein_aufzaehlungszeichen(para)

cp = p.core_properties
cp.title = "UML-Klassendiagramm - Vertiefung"
cp.subject = "MOBIS, Pruefungsleistung " + K.TERMIN
cp.author = ""
cp.last_modified_by = ""
cp.comments = ("Aufgabenfolien nach dem Format von Aufgabe 5 - Reklamation. "
               "Vorlage, Master und Theme stammen aus dem Kursmaterial.")
p.save(OUT)

print("geschrieben:", OUT)
for i, s in enumerate(p.slides):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text:
            art = "Titel" if sh.height < Cm(4) else "Text"
            print("  Folie %d  %-5s  %.1f x %.1f cm  |  %d Woerter"
                  % (i + 1, art, sh.width / 360000, sh.height / 360000,
                     len(sh.text_frame.text.split())))
