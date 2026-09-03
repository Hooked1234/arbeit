# -*- coding: utf-8 -*-
"""Strukturpruefung der erzeugten Dokumente ohne Office.

Word COM haengt in dieser Umgebung reproduzierbar beim Export, deshalb wird
hier ueber die Datei selbst geprueft: ZIP-Integritaet, Abschnittsformate,
Bilder, Tabellenzahl, Vollstaendigkeit der Pflichtinhalte.

Aufruf:  python _generator/check_dokumente.py <Dokumenteordner>
"""
import os
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from pptx import Presentation
from pptx.util import Cm

import content as K

D = sys.argv[1] if len(sys.argv) > 1 else ".."
LF = os.path.join(D, "MOBIS_UML_Vertiefung_Lehrendenfassung.docx")
AB = os.path.join(D, "MOBIS_UML_Vertiefung_Aufgabenblatt.docx")
PP = os.path.join(D, "MOBIS_UML_Vertiefung_Aufgaben.pptx")

errs, oks = [], []


def pruefe(bedingung, text):
    (oks if bedingung else errs).append(text)


def zip_ok(pfad):
    try:
        with zipfile.ZipFile(pfad) as z:
            return z.testzip() is None
    except Exception:
        return False


# ------------------------------------------------------------ Lehrendenfassung
pruefe(zip_ok(LF), "Lehrendenfassung: ZIP-Integritaet")
d = Document(LF)
volltext = "\n".join(p.text for p in d.paragraphs)
tab_text = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
alles = volltext + "\n" + tab_text

quer = [s for s in d.sections if s.orientation == WD_ORIENT.LANDSCAPE]
pruefe(len(quer) == 2,
       "Lehrendenfassung: zwei A3-Querseiten (gefunden %d)" % len(quer))
for s in quer:
    pruefe(abs(s.page_width - Cm(42.0)) < Cm(0.2),
           "Lehrendenfassung: Querseite ist A3 breit")

bilder = len(d.inline_shapes)
pruefe(bilder == 2,
       "Lehrendenfassung: zwei Diagramme eingebettet (gefunden %d)" % bilder)

for kap in ("1  Durchfuehrung", "2  Didaktische Leitplanken",
            "5  Betreuung zu zweit", "6  Abnahme vor der Veranstaltung"):
    pruefe(kap in volltext, "Lehrendenfassung: Kapitel '%s'" % kap)

pruefe(len(K.A1["rueck"]) == sum(
    1 for t in d.tables for r in t.rows
    if len(r.cells) == 2 and r.cells[0].text.strip() in
    [x[0] for x in K.A1["rueck"]]),
    "Lehrendenfassung: alle %d Rueckfuehrbarkeitszeilen vorhanden"
    % len(K.A1["rueck"]))

for nr, stelle, aussage, befund, korr in K.A2["fundliste"]:
    pruefe(stelle in tab_text, "Lehrendenfassung: Fund %s in der Fundliste" % nr)

fehlerzahl = sum(1 for f in K.A2["fundliste"] if f[3] == "Fehler")
zulaessig = sum(1 for f in K.A2["fundliste"] if f[3] == "zulaessig")
pruefe(fehlerzahl == 8 and zulaessig == 2,
       "Fundliste: 8 Fehler und 2 zulaessige Abweichungen "
       "(gefunden %d / %d)" % (fehlerzahl, zulaessig))

p1 = sum(int(r[1]) for r in K.A1["raster"])
p2 = sum(int(r[1]) for r in K.A2["raster"])
pruefe(p1 == 20, "Raster Aufgabe 1 ergibt 20 Punkte (ergibt %d)" % p1)
pruefe(p2 == 20, "Raster Aufgabe 2 ergibt 20 Punkte (ergibt %d)" % p2)

for frage, antwort in K.FAQ:
    pruefe(frage in tab_text, "Lehrendenfassung: FAQ '%s'" % frage[:34])

# ------------------------------------------------------------- Aufgabenblatt
pruefe(zip_ok(AB), "Aufgabenblatt: ZIP-Integritaet")
a = Document(AB)
a_text = "\n".join(p.text for p in a.paragraphs)
a_tab = "\n".join(c.text for t in a.tables for r in t.rows for c in r.cells)

pruefe(K.A1["szenario"].split("\n\n")[0] in a_text,
       "Aufgabenblatt: Szenariotext Aufgabe 1 wortgleich")
pruefe(K.A2["szenario"].split("\n\n")[0] in a_text,
       "Aufgabenblatt: Auftragstext Aufgabe 2 wortgleich")

protokoll = [t for t in a.tables if len(t.columns) == 5]
pruefe(len(protokoll) == 1, "Aufgabenblatt: genau ein Protokollbogen")
if protokoll:
    pruefe(len(protokoll[0].rows) == 11,
           "Aufgabenblatt: Protokollbogen mit 10 Zeilen plus Kopf "
           "(gefunden %d)" % (len(protokoll[0].rows) - 1))

# Zwei Stufen. Loesungswoerter benennen eine konkrete Modellierungsentscheidung
# und duerfen nirgends auf dem Blatt stehen. Notationswoerter sind Formvorgaben
# und im Arbeitsauftrag zulaessig - "Kardinalitaet" ist ausserdem das Wort des
# Dozenten von Folie 6. Im Szenariotext bleiben auch sie verboten.
LOESUNGSWOERTER = ["Aggregation", "Komposition", "Assoziation", "Vererbung",
                   "Generalisierung", "abstrakt", "Raute"]
NOTATIONSWOERTER = ["Kardinalit", "Multiplizit"]

gefunden = [w for w in LOESUNGSWOERTER if w in a_text or w in a_tab]
pruefe(not gefunden,
       "Aufgabenblatt: kein Loesungswort auf dem Blatt "
       "(gefunden: %s)" % ", ".join(gefunden))

for A in K.AUFGABEN:
    treffer = [w for w in LOESUNGSWOERTER + NOTATIONSWOERTER
               if w in A["szenario"]]
    pruefe(not treffer,
           "Szenariotext Aufgabe %d: kein Fachbegriff (gefunden: %s)"
           % (A["nr"], ", ".join(treffer)))

pruefe("Kardinalit" in a_text,
       "Aufgabenblatt: Arbeitsauftrag verlangt die Kardinalitaeten "
       "ausdruecklich")

# --------------------------------------------------------------------- PPTX
pruefe(zip_ok(PP), "Aufgabenfolien: ZIP-Integritaet")
p = Presentation(PP)
pruefe(len(p.slides) == 2,
       "Aufgabenfolien: genau zwei Folien (gefunden %d)" % len(p.slides))

for i, (slide, A) in enumerate(zip(p.slides, K.AUFGABEN)):
    txt = "\n".join(sh.text_frame.text for sh in slide.shapes
                    if sh.has_text_frame)
    pruefe(A["titel"] in txt, "Folie %d: Titel gesetzt" % (i + 1))
    pruefe(A["szenario"].split("\n\n")[0] in txt,
           "Folie %d: Text wortgleich zum Aufgabenblatt" % (i + 1))
    pruefe(not any(w in txt for w in LOESUNGSWOERTER),
           "Folie %d: kein Loesungswort" % (i + 1))

cp = p.core_properties
pruefe(not cp.author and not cp.last_modified_by,
       "Aufgabenfolien: Dokumenteigenschaften bereinigt")

# -------------------------------------------------------------------- Bericht
print("Strukturpruefung der Dokumente")
print("  bestanden: %d" % len(oks))
for e in errs:
    print("  FEHLER:", e)
if errs:
    print("MIT FEHLERN")
    raise SystemExit(1)
print("BESTANDEN: alle %d Pruefungen ohne Befund." % len(oks))
